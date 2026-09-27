# SPDX-FileCopyrightText: 2026 davlillos
# SPDX-License-Identifier: MIT

"""Convierte la Calificacion.xlsx del docente en la fixture de calibracion.

El docente califico 89 entregas a mano en 4 turnos, con una hoja por turno.
Cada bloque de alumno tiene la misma tabla:

    Criterio | % | Esperados | Modelados | Nota ponderada | Observacion

con la formula =IF(C>=D,(D/C)*B,IF(C<D,(C/D)*B,0)) y Total = SUMA(E)*10.

De ahi se extrae:
  - la rubrica de cada turno, leida con teacher_rubric_parser: la misma
    lectura que hace POST /api/rubric/import cuando el docente sube su Excel,
    asi la calibracion mide el camino real y no una copia de el
  - la nota final que puso el docente por alumno
  - su columna "Modelados", para poder comparar criterio por criterio

Los nombres de clase quedan como los escribio el docente ("Historial" por
"HistorialEnfermedad"): el comparador los empareja con la solucion igual que
cuando califica desde la pantalla.

Uso:
    python scripts/build_calibracion_fixture.py
"""
from __future__ import annotations

import json
import re
import sys
from dataclasses import asdict
from pathlib import Path
from typing import Any, List, Optional

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from app.parsers.teacher_rubric_parser import parse_teacher_worksheet  # noqa: E402

FIXTURE_DIR = BACKEND / "test_files" / "calibracion"
XLSX = FIXTURE_DIR / "calificacion_docente.xlsx"
OUT = FIXTURE_DIR / "rubricas.json"

# hoja del Excel -> (carpeta de entregas, xmi de solucion)
TURNOS = {
    "Tuno10AMImpar": ("T1-IMPAR", "Turno1Impar.xmi"),
    "Turno10AMPar": ("T1-PAR", "Turno1Par.xmi"),
    "Turno1110AMImpar": ("T2-IMPAR", "Turno2Ambas.xmi"),
    "Turno1110AMPar": ("T2-PAR", "Turno2Ambas.xmi"),
}

STUDENT_RE = re.compile(r"^[A-Z]{2}\d{5}$")


def _celda(value: Any) -> Optional[str]:
    """La columna Modelados tal como la escribio el docente ("5", "0", texto)."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def student_blocks(ws: Worksheet) -> List[dict]:
    """Parte la hoja en un bloque por alumno: carne, nota y columna Modelados.

    `ws` es la hoja con los valores calculados (data_only). Solo se guardan
    los Modelados de las filas que puntuan, en el mismo orden que las reglas.
    """
    found, current = [], None
    for row in ws.iter_rows(min_col=1, max_col=5, values_only=True):
        label = str(row[0]).strip() if row[0] is not None else ""
        if STUDENT_RE.match(label):
            current = {"carne": label, "modelados": [], "nota": None}
        elif label == "Total" and current is not None:
            nota = row[4]
            # el docente dejo alguna celda con error (#VALUE!): sin nota
            current["nota"] = round(float(nota), 4) if isinstance(nota, (int, float)) else None
            found.append(current)
            current = None
        elif current is not None and label and label != "Criterio":
            peso, esperados, modelados = row[1], row[2], row[3]
            if peso is not None and esperados is not None:
                current["modelados"].append(_celda(modelados))
    return found


def main() -> None:
    formulas = load_workbook(XLSX)
    valores = load_workbook(XLSX, data_only=True)
    salida = {}
    for hoja, (carpeta, solucion) in TURNOS.items():
        rubrica = parse_teacher_worksheet(formulas[hoja], valores[hoja])
        rules = [
            {k: v for k, v in asdict(rule).items() if v is not None}
            for rule in rubrica.rules
        ]
        total = sum(r["weight"] for r in rules)
        assert abs(total - 100.0) < 0.01, "%s: los pesos suman %s" % (hoja, total)

        bloques = student_blocks(valores[hoja])
        alumnos = {}
        for bloque in bloques:
            xmi = FIXTURE_DIR / carpeta / ("%s.xmi" % bloque["carne"])
            if bloque["nota"] is None or not xmi.exists():
                continue
            assert len(bloque["modelados"]) == len(rules), (
                "%s %s: %d filas que puntuan y %d reglas"
                % (hoja, bloque["carne"], len(bloque["modelados"]), len(rules))
            )
            alumnos[bloque["carne"]] = {
                "nota_docente": bloque["nota"],
                "modelados_docente": {
                    rule["rule_id"]: modelados
                    for rule, modelados in zip(rules, bloque["modelados"])
                },
            }
        salida[carpeta] = {
            "hoja_excel": hoja,
            "solucion": solucion,
            "reglas": rules,
            "alumnos": alumnos,
        }
        avisos = "" if not rubrica.warnings else " · avisos: %s" % rubrica.warnings
        print("%-10s %2d reglas, %2d alumnos con XMI y nota (de %d calificados)%s"
              % (carpeta, len(rules), len(alumnos), len(bloques), avisos))

    OUT.write_text(json.dumps(salida, ensure_ascii=False, indent=1), encoding="utf-8")
    print("\nEscrito %s" % OUT)


if __name__ == "__main__":
    main()

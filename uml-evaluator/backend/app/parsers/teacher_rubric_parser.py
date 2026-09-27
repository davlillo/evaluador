# SPDX-FileCopyrightText: 2026 davlillos
# SPDX-License-Identifier: MIT

"""Lee la rúbrica en el formato propio del docente (su hoja de Excel).

Es la tabla que él usa para calificar a mano:

    Criterio                              | %  | Esperados | Modelados | Nota ponderada | Observaciones
    Clases                                | 40 | 6         |           | =IF(C>=D, ...)  |
    Relaciones                            | =SUM(B7:B22)                                   <- sección, no puntúa
    Asociación Sala - Funcion                                                              <- grupo, no puntúa
      Multiplicidad 1 en Sala             | 5  | 1         |
      Multiplicidad 0..* en Funcion       | 5  | 1         |
    Asociación reflexiva Sala                                                              <- grupo reflexivo
      Multiplicidad 1 en Sala             | 5  | 1         |
      Multiplicidad 1 en Sala             | 5  | 1         |
    Clase de asociacion Obra - Actor      | 10 | 1         |
    Total                                 | =SUM(B5:B6)            | =SUM(E5:E22)/10

Solo las filas hoja puntúan (Clases, cada multiplicidad y cada clase de
asociación). Los encabezados de grupo dicen entre qué clases está la relación
y de qué tipo es; la fila "Relaciones" solo suma.

Se aceptan las dos variantes que usó el docente:

- 2EP (2026): pesos en porcentaje (40, 5, 10) que suman 100.
- Práctica 1 (año pasado): pesos en fracción (0.2, 0.05) que suman 1, y
  varios estudiantes por hoja; se toma el primer bloque.

La hoja se escribe a mano, así que el parser tolera tildes faltantes
("Composicion", "Clase de asociacion"), espacios alrededor del guion
("Sala - Funcion" o "Sala-Funcion") y typos en los nombres de clase de los
encabezados, que se corrigen contra los nombres de las filas de multiplicidad
("OdenServicio" -> "OrdenServicio") y se reportan como aviso.
"""
from __future__ import annotations

import difflib
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.comparator.scoring_modes import ClassRubricRule
from app.parsers.rubric_parser import RubricParseError

GROUP_RE = re.compile(
    r"^(?P<tipo>asociaci[oó]n\s+reflexiva|asociaci[oó]n|agregaci[oó]n|composici[oó]n)"
    r"\s+(?P<resto>.+)$",
    re.IGNORECASE,
)
MULTIPLICITY_RE = re.compile(
    r"^multiplicidad\s+(?P<valor>\S+)\s+en\s+(?P<clase>.+)$", re.IGNORECASE,
)
ASSOCIATION_CLASS_RE = re.compile(
    r"^clase\s+de\s+asociaci[oó]n\s+(?P<a>.+?)\s*-\s*(?P<b>.+)$", re.IGNORECASE,
)
PAIR_SEPARATOR_RE = re.compile(r"\s*-\s*")

#: Por debajo de esta similitud, un nombre del encabezado no se "corrige".
TYPO_CUTOFF = 0.8


@dataclass
class TeacherRubric:
    """Una rúbrica leída de la hoja del docente, lista para editar en pantalla."""

    title: str
    rules: List[ClassRubricRule]
    warnings: List[str] = field(default_factory=list)


@dataclass
class _Group:
    label: str
    relationship_type: str
    source: str
    target: str
    reflexive: bool
    row: int
    children: int = 0


def _norm(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", value or "")
    stripped = "".join(c for c in decomposed if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]", "", stripped.lower())


def _number(value: Any) -> Optional[float]:
    """Solo valores numéricos; una fórmula sin valor cacheado cuenta como vacía."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.replace(",", ".").strip())
        except ValueError:
            return None
    return None


def _multiplicity(value: str) -> str:
    """"n" y "N" se escriben a mano como sinónimo de "*"."""
    limpio = value.strip().replace(" ", "")
    return re.sub(r"(?<=\.\.)[nN]$|^[nN]$", "*", limpio)


def _relationship_type(tipo: str) -> str:
    base = _norm(tipo)
    if base.startswith("agregacion"):
        return "aggregation"
    if base.startswith("composicion"):
        return "composition"
    return "association"


def _find_header(ws: Worksheet) -> Optional[int]:
    for row in range(1, min(ws.max_row, 60) + 1):
        a = ws.cell(row=row, column=1).value
        b = ws.cell(row=row, column=2).value
        if isinstance(a, str) and _norm(a) == "criterio" and isinstance(b, str) and b.strip() == "%":
            return row
    return None


def _closest(name: str, candidates: List[str], cutoff: float) -> Optional[str]:
    by_norm = {_norm(c): c for c in candidates}
    match = difflib.get_close_matches(_norm(name), list(by_norm), n=1, cutoff=cutoff)
    return by_norm[match[0]] if match else None


def _canonicalize_groups(groups: List[_Group], known: List[str], warnings: List[str]) -> None:
    """Corrige typos de los encabezados contra los nombres de las filas hoja.

    El docente escribe el nombre de cada clase varias veces en la hoja; si en
    el encabezado puso "OdenServicio" pero en las multiplicidades
    "OrdenServicio", la versión correcta es la que se repite.
    """
    known_norm = {_norm(n) for n in known}
    for group in groups:
        for attr in ("source", "target"):
            name = getattr(group, attr)
            if _norm(name) in known_norm:
                continue
            fix = _closest(name, known, TYPO_CUTOFF)
            if fix:
                setattr(group, attr, fix)
                warnings.append(
                    f"En «{group.label}» se interpretó «{name}» como «{fix}»."
                )
    for group in groups:
        if group.reflexive:
            group.target = group.source


def _multiplicity_end(group: _Group, clase: str, row: int, errors: List[str]) -> Optional[str]:
    if group.reflexive:
        # los dos extremos son la misma clase: el orden de las filas decide
        return "source" if group.children == 0 else "target"
    if _norm(clase) == _norm(group.source):
        return "source"
    if _norm(clase) == _norm(group.target):
        return "target"
    fix = _closest(clase, [group.source, group.target], 0.6)
    if fix:
        return "source" if fix == group.source else "target"
    errors.append(
        f"Fila {row}: «{clase}» no es ninguno de los extremos de «{group.label}»."
    )
    return None


def parse_teacher_worksheet(ws: Worksheet, values: Worksheet) -> TeacherRubric:
    """Lee una hoja con el formato del docente.

    `ws` trae las fórmulas y `values` los valores cacheados de la misma hoja:
    los pesos de las filas hoja son números, pero no hay garantía.
    """
    header = _find_header(ws)
    if header is None:
        raise RubricParseError([
            f"La hoja «{ws.title}» no tiene la fila de encabezados "
            "(Criterio | % | Esperados | Modelados ...).",
        ])

    errors: List[str] = []
    warnings: List[str] = []
    pending: List[Dict[str, Any]] = []
    groups: List[_Group] = []
    group: Optional[_Group] = None
    known_names: List[str] = []

    for row in range(header + 1, ws.max_row + 1):
        raw_label = ws.cell(row=row, column=1).value
        if not isinstance(raw_label, str) or not raw_label.strip():
            continue
        label = " ".join(raw_label.split())
        if _norm(label) == "total":
            break  # el formato viejo trae más estudiantes debajo
        if _norm(label) == "relaciones":
            continue

        weight = _number(values.cell(row=row, column=2).value)
        if weight is None:
            weight = _number(ws.cell(row=row, column=2).value)
        expected = _number(values.cell(row=row, column=3).value)
        if expected is None:
            expected = _number(ws.cell(row=row, column=3).value)

        group_match = GROUP_RE.match(label)
        if group_match and expected is None:
            nombres = [n for n in PAIR_SEPARATOR_RE.split(group_match["resto"].strip()) if n]
            reflexive = "reflexiva" in _norm(group_match["tipo"]) or len(nombres) == 1
            if not nombres or len(nombres) > 2:
                errors.append(f"Fila {row}: no se entienden las clases de «{label}».")
                group = None
                continue
            group = _Group(
                label=label,
                relationship_type=_relationship_type(group_match["tipo"]),
                source=nombres[0],
                target=nombres[-1],
                reflexive=reflexive,
                row=row,
            )
            groups.append(group)
            continue

        if weight is None:
            errors.append(f"Fila {row}: «{label}» no tiene peso en la columna %.")
            continue

        if _norm(label) == "clases":
            pending.append({
                "criterion_type": "classes", "label": label, "weight": weight,
                "expected_quantity": int(expected) if expected is not None else None,
            })
            group = None
            continue

        mult = MULTIPLICITY_RE.match(label)
        if mult:
            if group is None:
                errors.append(
                    f"Fila {row}: «{label}» no cuelga de ningún encabezado de relación.",
                )
                continue
            clase = mult["clase"].strip()
            known_names.append(clase)
            pending.append({
                "criterion_type": "multiplicity", "label": label, "weight": weight,
                "group": group, "clase": clase, "row": row,
                "expected_multiplicity": _multiplicity(mult["valor"]),
            })
            continue

        assoc = ASSOCIATION_CLASS_RE.match(label)
        if assoc:
            known_names.extend([assoc["a"].strip(), assoc["b"].strip()])
            pending.append({
                "criterion_type": "association_class", "label": label, "weight": weight,
                "source": assoc["a"].strip(), "target": assoc["b"].strip(),
            })
            group = None
            continue

        errors.append(f"Fila {row}: no se reconoce el criterio «{label}».")

    if not pending and not errors:
        errors.append(f"La hoja «{ws.title}» no tiene criterios con peso.")
    if errors:
        raise RubricParseError(errors)

    _canonicalize_groups(groups, known_names, warnings)

    # los extremos se resuelven con los nombres ya corregidos, y en orden
    rules: List[ClassRubricRule] = []
    for index, item in enumerate(pending, start=1):
        common = {
            "rule_id": f"r{index}",
            "criterion_type": item["criterion_type"],
            "label": item["label"],
            "weight": item["weight"],
        }
        if item["criterion_type"] == "classes":
            rules.append(ClassRubricRule(**common, expected_quantity=item["expected_quantity"]))
        elif item["criterion_type"] == "multiplicity":
            g: _Group = item["group"]
            end = _multiplicity_end(g, item["clase"], item["row"], errors)
            g.children += 1
            rules.append(ClassRubricRule(
                **common,
                source=g.source,
                target=g.target,
                relationship_type=g.relationship_type,
                multiplicity_end=end,
                expected_multiplicity=item["expected_multiplicity"],
                group_label=g.label,
            ))
        else:
            rules.append(ClassRubricRule(
                **common,
                source=item["source"],
                target=item["target"],
                relationship_type="association_class",
            ))
    if errors:
        raise RubricParseError(errors)

    _scale_to_percent(rules, warnings)
    for group in groups:
        if group.children == 0:
            warnings.append(f"«{group.label}» no tiene filas de multiplicidad debajo.")
    return TeacherRubric(title=ws.title, rules=rules, warnings=warnings)


def _scale_to_percent(rules: List[ClassRubricRule], warnings: List[str]) -> None:
    """El formato viejo usa fracciones (0.2); el nuevo, porcentajes (40)."""
    total = sum(rule.weight for rule in rules)
    if abs(total - 1.0) < 0.01:
        for rule in rules:
            rule.weight = round(rule.weight * 100, 4)
        total = sum(rule.weight for rule in rules)
    if abs(total - 100.0) >= 0.01:
        warnings.append(
            f"Los pesos suman {total:g}% y deberían sumar 100%. Ajustalos en pantalla.",
        )


def is_teacher_sheet(ws: Worksheet) -> bool:
    return _find_header(ws) is not None


def parse_teacher_rubric_xlsx(file_path: str) -> TeacherRubric:
    """Lee la primera hoja del libro que tenga el formato del docente."""
    formulas = load_workbook(file_path)
    values = load_workbook(file_path, data_only=True)
    for ws in formulas.worksheets:
        if is_teacher_sheet(ws):
            return parse_teacher_worksheet(ws, values[ws.title])
    raise RubricParseError([
        "Ninguna hoja tiene la tabla de la rúbrica "
        "(una fila con Criterio | % | Esperados | Modelados).",
    ])
